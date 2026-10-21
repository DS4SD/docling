# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for text with Word's Hidden font property (``w:vanish``)."""

from pathlib import Path

from docling_core.types.doc import ContentLayer, TextItem
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docling.backend.msword_backend import MsWordDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import DoclingDocument, InputDocument


def _convert(document, tmp_path: Path) -> DoclingDocument:
    path = tmp_path / "hidden.docx"
    document.save(str(path))
    in_doc = InputDocument(
        path_or_stream=path, format=InputFormat.DOCX, backend=MsWordDocumentBackend
    )
    return in_doc._backend.convert()


def _layers_by_text(doc: DoclingDocument) -> dict[str, str]:
    return {
        item.text: item.content_layer.value
        for item, _ in doc.iterate_items(included_content_layers=set(ContentLayer))
        if isinstance(item, TextItem)
    }


def test_docx_hidden_text_is_left_out_of_the_default_export(tmp_path: Path):
    """Word does not show or print hidden text.

    A fully hidden paragraph goes to ``ContentLayer.INVISIBLE``, the same
    treatment a hidden PPTX slide or shape gets. A hidden run in a shown
    paragraph or a table cell cannot have its own layer, so it is dropped.
    """
    document = Document()
    para = document.add_paragraph("Shown start. ")
    para.add_run("Hidden run. ").font.hidden = True
    para.add_run("Shown end.")
    document.add_paragraph().add_run("Hidden paragraph.").font.hidden = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Shown cell"
    table.cell(0, 1).paragraphs[0].add_run("Hidden cell").font.hidden = True

    doc = _convert(document, tmp_path)

    layers = _layers_by_text(doc)
    assert layers["Shown start. Shown end."] == "body"
    assert layers["Hidden paragraph."] == "invisible"
    assert [c.text for c in doc.tables[0].data.table_cells] == ["Shown cell", ""]
    assert "Hidden" not in doc.export_to_markdown()


def test_docx_hidden_text_follows_style_inheritance(tmp_path: Path):
    """Hidden can come from a character or paragraph style, and an explicit
    ``w:vanish w:val="0"`` on a run overrides a hidden style."""
    document = Document()
    char_style = document.styles.add_style("Hidden Chars", WD_STYLE_TYPE.CHARACTER)
    char_style.font.hidden = True
    para_style = document.styles.add_style("Hidden Para", WD_STYLE_TYPE.PARAGRAPH)
    para_style.font.hidden = True

    document.add_paragraph("Shown. ").add_run("Hidden by character style.", char_style)
    document.add_paragraph("Hidden by paragraph style.", para_style)
    mixed = document.add_paragraph(style=para_style)
    mixed.add_run("Hidden by paragraph style too. ")
    mixed.add_run("Shown by override.").font.hidden = False

    doc = _convert(document, tmp_path)

    layers = _layers_by_text(doc)
    assert layers["Shown."] == "body"
    assert layers["Hidden by paragraph style."] == "invisible"
    assert layers["Shown by override."] == "body"
    assert "Hidden" not in doc.export_to_markdown()
