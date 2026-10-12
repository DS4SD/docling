# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for PowerPoint combination charts: several plots in one chart."""

import datetime
import logging
from pathlib import Path

from docling_core.types.doc import PictureClassificationLabel, PictureItem

from .test_backend_pptx import get_converter

_MONTHS = [datetime.date(2012, month, 1) for month in range(1, 13)]
_BARS = {
    "East": (85, 34, 100, 80, 116, 88, 71, 90, 121, 110, 101, 81),
    "West": (94, 72, 62, 91, 99, 70, 71, 46, 91, 57, 32, 157),
    "North": (91, 59, 85, 67, 130, 66, 51, 40, 62, 61, 61, 52),
    "South": (71, 110, 101, 28, 39, 52, 48, 125, 43, 122, 60, 63),
}
_LINES = {
    "Margin": tuple(
        n / 1000 for n in (126, 74, 110, 92, 86, 89, 77, 40, 101, 62, 81, 62)
    ),
    "Target": tuple(
        n / 1000 for n in (91, 75, 94, 72, 103, 75, 66, 81, 85, 94, 69, 95)
    ),
}


def _save_combo_chart_deck(path: Path, second_plot: str = "lineChart") -> None:
    """Save a deck whose chart draws columns and a second plot on a second axis.

    python-pptx writes single-plot charts only, so the chart starts as clustered
    columns over date categories. The last two series then move into a
    ``c:<second_plot>`` element on a secondary value axis, which is how
    PowerPoint stores a "Clustered Column - Line on Secondary Axis" chart. With
    ``second_plot="stockChart"`` the second plot is a stock plot, as in a
    "Volume-High-Low-Close" chart, for which python-pptx has no plot class.
    """
    import copy

    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    chart_data = CategoryChartData()
    chart_data.categories = _MONTHS
    for name, values in {**_BARS, **_LINES}.items():
        chart_data.add_series(name, values)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(0.5),
        Inches(9),
        Inches(6),
        chart_data,
    ).chart

    plot_area = chart.element.chart.plotArea
    bar_chart = plot_area.find(qn("c:barChart"))
    second = etree.Element(qn(f"c:{second_plot}"))
    if second_plot == "lineChart":
        etree.SubElement(second, qn("c:grouping")).set("val", "standard")
        etree.SubElement(second, qn("c:varyColors")).set("val", "0")
    for ser in bar_chart.findall(qn("c:ser"))[len(_BARS) :]:
        for invert in ser.findall(qn("c:invertIfNegative")):
            ser.remove(invert)
        second.append(ser)
    if second_plot == "lineChart":
        etree.SubElement(second, qn("c:marker")).set("val", "1")
    else:
        etree.SubElement(second, qn("c:hiLowLines"))
    etree.SubElement(second, qn("c:axId")).set("val", "900")
    etree.SubElement(second, qn("c:axId")).set("val", "901")
    bar_chart.addnext(second)

    category_axis = plot_area.find(qn("c:dateAx"))
    value_axis = plot_area.find(qn("c:valAx"))
    secondary_category_axis = copy.deepcopy(category_axis)
    secondary_category_axis.find(qn("c:axId")).set("val", "900")
    secondary_category_axis.find(qn("c:delete")).set("val", "1")
    secondary_category_axis.find(qn("c:crossAx")).set("val", "901")
    secondary_value_axis = copy.deepcopy(value_axis)
    secondary_value_axis.find(qn("c:axId")).set("val", "901")
    secondary_value_axis.find(qn("c:crossAx")).set("val", "900")
    value_axis.addnext(secondary_value_axis)
    secondary_value_axis.addnext(secondary_category_axis)
    prs.save(path)


def _only_chart(path: Path) -> PictureItem:
    doc = get_converter().convert(path).document
    pictures = list(doc.pictures)
    assert len(pictures) == 1
    return pictures[0]


def _grid(picture: PictureItem) -> list[list[str]]:
    assert picture.meta is not None and picture.meta.tabular_chart is not None
    return [
        [cell.text for cell in row]
        for row in picture.meta.tabular_chart.chart_data.grid
    ]


def _labels(picture: PictureItem) -> list[str]:
    assert picture.meta is not None and picture.meta.classification is not None
    return [p.class_name for p in picture.meta.classification.predictions]


def test_combo_chart_lists_each_plot_type_and_keeps_all_series(tmp_path: Path):
    """A column + line chart is classified as both, back-most plot first.

    The main prediction stays the column plot, as before, so serializers that
    read only the main prediction still treat the picture as a bar chart. The
    line series on the secondary axis stay in the data table, after the column
    series.
    """
    path = tmp_path / "combo_chart.pptx"
    _save_combo_chart_deck(path)

    picture = _only_chart(path)

    assert _labels(picture) == [
        PictureClassificationLabel.BAR_CHART.value,
        PictureClassificationLabel.LINE_CHART.value,
    ]
    grid = _grid(picture)
    assert grid[0][1:] == ["East", "West", "North", "South", "Margin", "Target"]
    assert grid[1][1:] == ["85", "94", "91", "71", "0.126", "0.091"]
    assert len(grid) == 1 + len(_MONTHS)


def test_unreadable_plot_keeps_the_data_of_the_other_plots(tmp_path: Path, caplog):
    """A plot type python-pptx cannot read does not drop the other plots' data.

    python-pptx has no plot class for a stock plot. Before, the whole chart
    lost its data table. Now the column series stay, the stock plot counts as
    another chart type, and a warning names the plot that was skipped.
    """
    path = tmp_path / "stock_combo_chart.pptx"
    _save_combo_chart_deck(path, second_plot="stockChart")

    with caplog.at_level(
        logging.WARNING, logger="docling.backend.mspowerpoint_backend"
    ):
        picture = _only_chart(path)

    assert _labels(picture) == [
        PictureClassificationLabel.BAR_CHART.value,
        PictureClassificationLabel.OTHER_CHART.value,
    ]
    assert _grid(picture)[0][1:] == list(_BARS)
    assert "stockChart" in caplog.text
